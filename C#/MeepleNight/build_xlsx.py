import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
RAW = open(os.path.join(ROOT, 'match_details.json'), encoding='utf-8').read()
matches_by_div = json.loads(json.loads(RAW))

OUT = os.path.expanduser(r'~\Documents\IVL_Season6_TopStats.xlsx')

# Re-aggregate per division, now including errors and team totals
data = {}
for div in ['A', 'B', 'C']:
    players = {}
    teams = {}
    team_matches = {}  # team -> set of match ids
    for m in matches_by_div[div]:
        for team in m['teams']:
            tname = team['teamName']
            team_matches.setdefault(tname, set()).add(m['id'])
            if tname not in teams:
                teams[tname] = dict(team=tname, matches=0, ace=0, attack=0, block=0,
                                    serviceError=0, foul=0, error=0)
            for p in team['players']:
                k = p['pid']
                if k not in players:
                    players[k] = dict(name=p['name'], team=tname, matches=0,
                                      ace=0, attack=0, block=0,
                                      serviceError=0, foul=0, error=0)
                e = players[k]
                e['matches'] += 1
                e['ace'] += p['ace']
                e['attack'] += p['attack']
                e['block'] += p['block']
                e['serviceError'] += p['serviceError']
                e['foul'] += p['foul']
                e['error'] += p['error']
                e['team'] = tname
                t = teams[tname]
                t['ace'] += p['ace']
                t['attack'] += p['attack']
                t['block'] += p['block']
                t['serviceError'] += p['serviceError']
                t['foul'] += p['foul']
                t['error'] += p['error']
    for tname, t in teams.items():
        t['matches'] = len(team_matches[tname])
        t['points'] = t['ace'] + t['attack'] + t['block']
        t['totalErrors'] = t['serviceError'] + t['foul'] + t['error']
    lst = []
    for pid, v in players.items():
        v['pid'] = pid
        v['points'] = v['ace'] + v['attack'] + v['block']
        v['totalErrors'] = v['serviceError'] + v['foul'] + v['error']
        lst.append(v)
    def topBy(key, n=5):
        return sorted(lst, key=lambda x: (-x[key], -x['points']))[:n]
    data[div] = {
        'matchCount': len(matches_by_div[div]),
        'playerCount': len(lst),
        'topPoints':   topBy('points'),
        'topAttack':   topBy('attack'),
        'topBlock':    topBy('block'),
        'topAce':      topBy('ace'),
        'topErrors':   topBy('totalErrors'),
        'allPlayers':  sorted(lst, key=lambda x: -x['points']),
        'teams':       sorted(teams.values(), key=lambda x: -x['points']),
    }

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill('solid', fgColor='1F4E78')
err_fill    = PatternFill('solid', fgColor='C00000')
header_font = Font(bold=True, color='FFFFFF', size=11)
section_font = Font(bold=True, size=13, color='1F4E78')
section_err_font = Font(bold=True, size=13, color='C00000')
center = Alignment(horizontal='center', vertical='center')
thin = Side(border_style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

DIV_NAMES = {'A': 'Мъже А Дивизия', 'B': 'Мъже Б Дивизия', 'C': 'Мъже В Дивизия'}

def write_top_table(ws, start_row, title, rows, cols, *, error_table=False):
    title_cell = ws.cell(row=start_row, column=1, value=title)
    title_cell.font = section_err_font if error_table else section_font
    r = start_row + 1
    headers = ['#'] + [c[0] for c in cols]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = header_font
        c.fill = err_fill if error_table else header_fill
        c.alignment = center
        c.border = border
    r += 1
    for i, row in enumerate(rows, 1):
        ws.cell(row=r, column=1, value=i).border = border
        ws.cell(row=r, column=1).alignment = center
        for ci, (_, key) in enumerate(cols, 2):
            cell = ws.cell(row=r, column=ci, value=row.get(key, ''))
            cell.border = border
            if isinstance(row.get(key), (int, float)):
                cell.alignment = center
        r += 1
    return r + 1

for div in ['A', 'B', 'C']:
    d = data[div]
    ws = wb.create_sheet(f'Дивизия {div}')
    ws.cell(row=1, column=1, value=f"{DIV_NAMES[div]} — Сезон 6 (Сезон 2026)").font = Font(bold=True, size=16, color='1F4E78')
    ws.cell(row=2, column=1, value=f"Мачове: {d['matchCount']}    Играчи: {d['playerCount']}").font = Font(italic=True, color='595959')

    row = 4
    row = write_top_table(ws, row, 'Топ 5 — Най-много точки (ace + attack + block)', d['topPoints'],
                          [('Играч', 'name'), ('Отбор', 'team'), ('Мачове', 'matches'),
                           ('Точки', 'points'), ('Атаки', 'attack'), ('Блокади', 'block'), ('Аса', 'ace'),
                           ('Грешки', 'totalErrors')])
    row = write_top_table(ws, row, 'Топ 5 — Атаки', d['topAttack'],
                          [('Играч', 'name'), ('Отбор', 'team'), ('Мачове', 'matches'), ('Атаки', 'attack')])
    row = write_top_table(ws, row, 'Топ 5 — Блокади', d['topBlock'],
                          [('Играч', 'name'), ('Отбор', 'team'), ('Мачове', 'matches'), ('Блокади', 'block')])
    row = write_top_table(ws, row, 'Топ 5 — Аса', d['topAce'],
                          [('Играч', 'name'), ('Отбор', 'team'), ('Мачове', 'matches'), ('Аса', 'ace')])
    row = write_top_table(ws, row, 'Топ 5 — Най-много грешки (общо)', d['topErrors'],
                          [('Играч', 'name'), ('Отбор', 'team'), ('Мачове', 'matches'),
                           ('Грешки общо', 'totalErrors'), ('Грешки сервис', 'serviceError'),
                           ('Нарушения', 'foul'), ('Други грешки', 'error')],
                          error_table=True)

    # Team totals table
    row += 1
    ws.cell(row=row, column=1, value='Класиране на отборите — точки и грешки общо').font = section_font
    row += 1
    team_headers = ['#', 'Отбор', 'Мачове', 'Точки', 'Атаки', 'Блокади', 'Аса',
                    'Грешки общо', 'Гр.сервис', 'Нарушения', 'Др.грешки']
    team_header_row = row
    for ci, h in enumerate(team_headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = center; c.border = border
    row += 1
    for i, t in enumerate(d['teams'], 1):
        vals = [i, t['team'], t['matches'], t['points'], t['attack'], t['block'], t['ace'],
                t['totalErrors'], t['serviceError'], t['foul'], t['error']]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=ci, value=v)
            cell.border = border
            if ci != 2:
                cell.alignment = center
        row += 1
    team_data_end = row - 1
    last_team_col = get_column_letter(len(team_headers))

    widths = {1: 5, 2: 28, 3: 9, 4: 9, 5: 9, 6: 10, 7: 8, 8: 12, 9: 11, 10: 12, 11: 12}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = 'A4'

    # Separate sheet: full per-player ranking with autofilter
    ws2 = wb.create_sheet(f'Дивизия {div} — Играчи')
    ws2.cell(row=1, column=1, value=f"{DIV_NAMES[div]} — Сезон 6 — Всички играчи").font = Font(bold=True, size=16, color='1F4E78')
    ws2.cell(row=2, column=1, value=f"Играчи: {d['playerCount']}    (използвайте филтрите за сортиране)").font = Font(italic=True, color='595959')
    full_headers = ['#', 'Играч', 'Отбор', 'Мачове', 'Точки', 'Атаки', 'Блокади', 'Аса',
                    'Грешки общо', 'Гр.сервис', 'Нарушения', 'Др.грешки']
    full_header_row = 4
    for ci, h in enumerate(full_headers, 1):
        c = ws2.cell(row=full_header_row, column=ci, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = center; c.border = border
    r2 = full_header_row + 1
    for i, p in enumerate(d['allPlayers'], 1):
        vals = [i, p['name'], p['team'], p['matches'], p['points'], p['attack'], p['block'], p['ace'],
                p['totalErrors'], p['serviceError'], p['foul'], p['error']]
        for ci, v in enumerate(vals, 1):
            cell = ws2.cell(row=r2, column=ci, value=v)
            cell.border = border
            if ci != 2 and ci != 3:
                cell.alignment = center
        r2 += 1
    full_data_end = r2 - 1
    last_col_letter = get_column_letter(len(full_headers))
    ws2.auto_filter.ref = f'A{full_header_row}:{last_col_letter}{full_data_end}'
    widths2 = {1: 5, 2: 28, 3: 26, 4: 9, 5: 9, 6: 9, 7: 10, 8: 8, 9: 12, 10: 11, 11: 12, 12: 12}
    for col, w in widths2.items():
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.freeze_panes = f'A{full_header_row + 1}'

ws = wb.create_sheet('Източник', 0)
ws.cell(row=1, column=1, value='IVL София — Сезон 6 — Мъже — Топ играчи').font = Font(bold=True, size=18, color='1F4E78')
ws.cell(row=3, column=1, value='Източник:').font = Font(bold=True)
ws.cell(row=3, column=2, value='ivl.bg / sportzonelive.azurewebsites.net (бекенд)')
ws.cell(row=4, column=1, value='Период:').font = Font(bold=True)
ws.cell(row=4, column=2, value='Сезон 6 — януари–април 2026')
ws.cell(row=5, column=1, value='"Точки" =').font = Font(bold=True)
ws.cell(row=5, column=2, value='ace + attack + block (трите видове точкови действия)')
ws.cell(row=6, column=1, value='"Грешки общо" =').font = Font(bold=True)
ws.cell(row=6, column=2, value='serviceError (грешки от сервис) + foul (нарушения) + error (други грешки в игра)')
ws.cell(row=8, column=1, value='Бележки:').font = Font(bold=True)
ws.cell(row=9, column=2, value='• Дивизия А: 20 мача (5 отбора, всеки играе с всеки 2 пъти)')
ws.cell(row=10, column=2, value='• Дивизия Б: 29 мача в API (стандартно 28: 8 отбора × 7 мача) — 1 допълнителен мач включен')
ws.cell(row=11, column=2, value='• Дивизия В: 30 мача (6 отбора, всеки играе с всеки 2 пъти)')
ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 75

wb.save(OUT)
print(f'Saved: {OUT}')

